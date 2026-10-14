"""Compare two machine lists and produce deterministic, auditable reports."""
import argparse
import csv
import hashlib
import json
import unicodedata
from collections import defaultdict
from pathlib import Path

VERSION = "1.0.0"


def normalize(value):
    """Normalize Unicode, case and whitespace without discarding model punctuation."""
    return " ".join(unicodedata.normalize("NFKC", str(value)).split()).casefold()


def read_lists(path, sheet=None, header=False):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            ws = book[sheet] if sheet else book.worksheets[0]
            selected = ws.title
            rows = []
            for cells in ws.iter_rows(min_col=1, max_col=2):
                if any(c.data_type in ("f", "e") for c in cells):
                    raise ValueError("Formula/error cells in A:B must be replaced with values")
                rows.append([c.value for c in cells])
        finally:
            book.close()
    elif path.suffix.lower() == ".csv":
        if sheet:
            raise ValueError("--sheet applies only to XLSX files")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        selected = None
    else:
        raise ValueError("Input must be .xlsx or .csv")
    lists = [defaultdict(list), defaultdict(list)]
    for number, row in enumerate(rows, 1):
        if header and number == 1:
            continue
        for col in range(2):
            value = row[col] if col < len(row) else None
            if value is None or not normalize(value):
                continue
            lists[col][normalize(value)].append({"row": number, "value": str(value)})
    return lists, selected


def compare(path, output, sheet=None, header=False):
    path, output = Path(path), Path(output)
    source_bytes = path.read_bytes()
    lists, selected = read_lists(path, sheet, header)
    a, b = lists
    keys = sorted(set(a) | set(b))
    records = []
    for key in keys:
        records.append({"key": key, "status": "both" if key in a and key in b else "only_a" if key in a else "only_b",
                        "a": a.get(key, []), "b": b.get(key, [])})
    summary = {"version": VERSION, "source_name": path.name,
               "source_sha256": hashlib.sha256(source_bytes).hexdigest(),
               "sheet": selected, "columns": ["A", "B"], "skip_header": header,
               "matching": "Unicode NFKC, whitespace collapse, casefold; punctuation retained",
               "counts": {"rows_a": sum(map(len, a.values())), "rows_b": sum(map(len, b.values())),
                          "unique_a": len(a), "unique_b": len(b),
                          "both": len(set(a) & set(b)), "only_a": len(set(a) - set(b)),
                          "only_b": len(set(b) - set(a)),
                          "duplicate_keys_a": sum(len(v) > 1 for v in a.values()),
                          "duplicate_keys_b": sum(len(v) > 1 for v in b.values())}}
    output.mkdir(parents=True, exist_ok=True)
    payload = {"summary": summary, "machines": records}
    (output / "comparison.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for status, filename in [("both", "matches.csv"), ("only_a", "only_a.csv"), ("only_b", "only_b.csv")]:
        with (output / filename).open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            writer.writerow(["normalized_name", "a_occurrences_json", "b_occurrences_json"])
            for record in records:
                if record["status"] == status:
                    # JSON strings preserve names and keep spreadsheet formula prefixes quoted.
                    writer.writerow([json.dumps(record["key"], ensure_ascii=False),
                                     json.dumps(record["a"], ensure_ascii=False), json.dumps(record["b"], ensure_ascii=False)])
    lines = ["# Equipment comparison report", "", f"Source: {path.name}",
             f"SHA-256: `{summary['source_sha256']}`", "", "| Metric | Count |", "| --- | ---: |"]
    lines.extend(f"| {key} | {value} |" for key, value in summary["counts"].items())
    lines.extend(["", "Counts represent unique normalized names unless labeled rows.",
                  "See matches.csv for shared machines and source row references.",
                  "Different wording, abbreviations and model variants require manual review.", ""])
    (output / "summary.md").write_text("\n".join(lines), encoding="utf-8")
    return payload


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, default=Path("reports/local"))
    parser.add_argument("--sheet", help="XLSX sheet name; default: first sheet")
    parser.add_argument("--header", action="store_true", help="Skip first row (original workbook has no header)")
    args = parser.parse_args()
    try:
        if args.input.resolve().parent == args.output.resolve() and args.input.name in {"comparison.json", "matches.csv", "only_a.csv", "only_b.csv", "summary.md"}:
            raise ValueError("Output would overwrite input")
        result = compare(args.input, args.output, args.sheet, args.header)
    except (OSError, ValueError, KeyError) as exc:
        parser.exit(2, f"Error: {exc}\n")
    print(json.dumps(result["summary"]["counts"], sort_keys=True))


if __name__ == "__main__":
    main()
